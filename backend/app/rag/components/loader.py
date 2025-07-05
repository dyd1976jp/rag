"""
文档加载器组件

支持多种文档格式的加载，包括PDF、TXT、MD等。
实现IDocumentLoader接口，提供统一的文档加载功能。
"""

import os
import uuid
import logging
from pathlib import Path
from typing import Union, Dict, Any, List

from app.rag.interfaces import BaseDocumentLoader, Document
from app.core.config import settings

logger = logging.getLogger(__name__)

class DocumentLoader(BaseDocumentLoader):
    """文档加载器实现"""
    
    def __init__(
        self,
        supported_formats: List[str] = None,
        max_file_size: int = None
    ):
        """
        初始化文档加载器
        
        Args:
            supported_formats: 支持的文件格式列表
            max_file_size: 最大文件大小（字节）
        """
        self.supported_formats = supported_formats or [
            '.pdf', '.txt', '.md', '.doc', '.docx'
        ]
        self.max_file_size = max_file_size or (50 * 1024 * 1024)  # 50MB
        
    async def load_document(self, file_path: Union[str, Path]) -> Document:
        """
        加载文档
        
        Args:
            file_path: 文档文件路径
            
        Returns:
            Document: 加载的文档对象
            
        Raises:
            ValueError: 不支持的文件格式或文件过大
            FileNotFoundError: 文件不存在
        """
        file_path = Path(file_path)
        
        # 检查文件是否存在
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        # 检查文件格式
        if not self.supports_format(file_path.suffix.lower()):
            raise ValueError(f"不支持的文件格式: {file_path.suffix}")
        
        # 检查文件大小
        file_size = file_path.stat().st_size
        if file_size > self.max_file_size:
            raise ValueError(f"文件过大: {file_size} bytes, 最大支持: {self.max_file_size} bytes")
        
        # 根据文件类型加载内容
        content = await self._load_content(file_path)
        
        # 创建文档对象
        document = Document(
            content=content,
            metadata={
                "file_name": file_path.name,
                "file_path": str(file_path),
                "file_size": file_size,
                "file_extension": file_path.suffix.lower(),
                "doc_id": str(uuid.uuid4())
            },
            doc_id=str(uuid.uuid4()),
            file_path=str(file_path)
        )
        
        logger.info(f"成功加载文档: {file_path.name}, 大小: {len(content)} 字符")
        return document
    
    def supports_format(self, file_extension: str) -> bool:
        """
        检查是否支持指定格式
        
        Args:
            file_extension: 文件扩展名
            
        Returns:
            bool: 是否支持该格式
        """
        return file_extension.lower() in self.supported_formats
    
    async def _load_content(self, file_path: Path) -> str:
        """
        根据文件类型加载内容
        
        Args:
            file_path: 文件路径
            
        Returns:
            str: 文件内容
        """
        extension = file_path.suffix.lower()
        
        if extension == '.pdf':
            return await self._load_pdf(file_path)
        elif extension in ['.txt', '.md']:
            return await self._load_text(file_path)
        elif extension in ['.doc', '.docx']:
            return await self._load_word(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {extension}")
    
    async def _load_pdf(self, file_path: Path) -> str:
        """加载PDF文件"""
        try:
            import PyPDF2
            
            content = ""
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    content += page.extract_text() + "\n"
            
            return content.strip()
        except ImportError:
            raise ImportError("需要安装PyPDF2库来处理PDF文件")
        except Exception as e:
            logger.error(f"PDF文件加载失败: {e}")
            raise ValueError(f"PDF文件加载失败: {e}")
    
    async def _load_text(self, file_path: Path) -> str:
        """加载文本文件"""
        try:
            # 尝试多种编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        return file.read()
                except UnicodeDecodeError:
                    continue
            
            raise ValueError("无法解码文件，尝试了多种编码格式")
            
        except Exception as e:
            logger.error(f"文本文件加载失败: {e}")
            raise ValueError(f"文本文件加载失败: {e}")
    
    async def _load_word(self, file_path: Path) -> str:
        """加载Word文件"""
        try:
            import docx
            
            doc = docx.Document(file_path)
            content = ""
            for paragraph in doc.paragraphs:
                content += paragraph.text + "\n"
            
            return content.strip()
        except ImportError:
            raise ImportError("需要安装python-docx库来处理Word文件")
        except Exception as e:
            logger.error(f"Word文件加载失败: {e}")
            raise ValueError(f"Word文件加载失败: {e}")
    
    def get_supported_formats(self) -> List[str]:
        """获取支持的文件格式列表"""
        return self.supported_formats.copy()
    
    def set_max_file_size(self, size: int) -> None:
        """设置最大文件大小"""
        self.max_file_size = size
        
    def add_supported_format(self, extension: str) -> None:
        """添加支持的文件格式"""
        if extension.lower() not in self.supported_formats:
            self.supported_formats.append(extension.lower())
